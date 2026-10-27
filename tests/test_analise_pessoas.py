"""Fase 15: os relatórios de pessoas da Análise. Presenças e ausências,
horas por área, e atrasos e faltas de marcação face aos turnos publicados,
só para o Admin e o Gerente."""
import datetime
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from analytics import people_reports as pr
from attendance.month import month_range, previous_month
from attendance.services import close_shift, open_shift
from core.models import Restaurant, WorkArea
from shifts import services as turnos
from shifts.models import Shift
from vacations.models import AbsenceType
from vacations.services import create_request, decide_request
from workers.models import Worker

pytestmark = pytest.mark.django_db
MES = previous_month(timezone.localdate())
INICIO, FIM = month_range(MES)
T = datetime.time


def _quando(dia, h, m=0):
    return timezone.make_aware(datetime.datetime.combine(dia, T(h, m)))


def _turno(casa, worker, dia, inicio, fim):
    s = turnos.create_shift(restaurant=casa, worker=worker, day=dia, start=T(inicio), end=T(fim))
    s.state = Shift.State.PUBLICADO
    s.save()
    return s


def _ponto(worker, dia, entrada, saida, casa=None, area=None):
    open_shift(worker=worker, at=_quando(dia, *entrada), restaurant=casa, work_area=area)
    close_shift(worker=worker, at=_quando(dia, *saida))


@pytest.fixture
def cenario(restaurant, manager):
    casa_b = Restaurant.objects.create(organization=restaurant.organization, name="Tasca da Baixa", slug="tasca-da-baixa")
    cozinha = WorkArea.objects.get(restaurant=restaurant, name="Cozinha")
    sala = WorkArea.objects.get(restaurant=restaurant, name="Sala")
    marta = Worker.objects.create(restaurant=restaurant, full_name="Marta Silva")
    rui = Worker.objects.create(restaurant=restaurant, full_name="Rui Costa")
    ana = Worker.objects.create(restaurant=casa_b, full_name="Ana da Baixa")
    ana.other_restaurants.add(restaurant)
    d = [INICIO + datetime.timedelta(days=i) for i in range(2, 8)]
    # Atraso de 20 min.
    _turno(restaurant, marta, d[0], 10, 16)
    _ponto(marta, d[0], (10, 20), (16, 0), area=cozinha)
    # Falta de marcação.
    _turno(restaurant, marta, d[1], 12, 16)
    # Saída uma hora antes.
    _turno(restaurant, marta, d[2], 10, 16)
    _ponto(marta, d[2], (10, 0), (15, 0), area=cozinha)
    # Sem turno: só horas.
    _ponto(marta, d[3], (9, 0), (13, 0), area=sala)
    # Dentro da tolerância: nada.
    _turno(restaurant, rui, d[0], 18, 23)
    _ponto(rui, d[0], (18, 5), (23, 0), area=sala)
    # Quem vem de outra casa conta nas horas desta.
    _ponto(ana, d[4], (12, 0), (16, 0), casa=restaurant, area=sala)
    # Uma entrada sem saída, por resolver: conta como presença, sem horas.
    open_shift(worker=rui, at=_quando(d[5], 18))
    # Férias do Rui, aprovadas: três dias úteis inteiros no mês.
    ferias = AbsenceType.objects.get(restaurant=restaurant, name="Férias")
    seg = INICIO + datetime.timedelta(days=(7 - INICIO.weekday()) % 7 + 7)
    pedido = create_request(worker=rui, absence_type=ferias, start_date=seg, end_date=seg + datetime.timedelta(days=2))
    decide_request(request_obj=pedido, user=manager, approve=True)
    return {"marta": marta, "rui": rui, "ana": ana, "dias": d, "casa_b": casa_b}


class TestContas:
    def test_presencas_e_desvios(self, restaurant, cenario):
        linhas, desvios = pr.pessoas(restaurant, INICIO, FIM)
        por = {p.worker.full_name: p for p in linhas}
        marta, rui, ana = por["Marta Silva"], por["Rui Costa"], por["Ana da Baixa"]
        assert marta.dias_trabalhados == 3
        assert marta.horas == datetime.timedelta(hours=5, minutes=40) + datetime.timedelta(hours=5) + datetime.timedelta(hours=4)
        assert marta.previstas == datetime.timedelta(hours=16)
        assert (marta.atrasos, marta.minutos_atraso, marta.saidas_cedo, marta.faltas) == (1, 20, 1, 1)
        assert (rui.atrasos, rui.faltas, rui.por_resolver, rui.dias_trabalhados) == (0, 0, 1, 2)
        assert rui.ausencias == Decimal(3)
        assert ana.visita and ana.horas == datetime.timedelta(hours=4)
        assert [(d.worker.full_name, d.tipo, d.minutos) for d in desvios] == [
            ("Marta Silva", "Atraso", 20),
            ("Marta Silva", "Falta de marcação", None),
            ("Marta Silva", "Saída antes da hora", 60),
        ]

    def test_horas_por_area(self, restaurant, cenario):
        areas = {a["area"]: a for a in pr.horas_por_area(restaurant, INICIO, FIM)}
        assert areas["Cozinha"]["horas"] == datetime.timedelta(hours=10, minutes=40)
        assert areas["Sala"]["horas"] == datetime.timedelta(hours=4 + 4 + 4, minutes=55)
        assert areas["Sala"]["pessoas"] == 3

    def test_ausencias_por_tipo_e_atravessar_o_mes(self, restaurant, manager, cenario):
        assert pr.ausencias_por_tipo(restaurant, INICIO, FIM) == [{"tipo": "Férias", "pedidos": 1, "dias": Decimal(3), "pessoas": 1}]
        # Uma ausência que começa no mês anterior conta só os dias deste.
        ferias = AbsenceType.objects.get(restaurant=restaurant, name="Férias")
        from vacations.services import count_days

        pedido = create_request(worker=cenario["marta"], absence_type=ferias, start_date=INICIO - datetime.timedelta(days=6),
                                end_date=INICIO + datetime.timedelta(days=1))
        decide_request(request_obj=pedido, user=manager, approve=True)
        dias = pr._dias_no_periodo(pedido, INICIO, FIM)
        assert dias == count_days(cenario["marta"], INICIO, INICIO + datetime.timedelta(days=1))
        assert dias < pedido.days_taken

    def test_presencas_por_mes(self, restaurant, cenario):
        meses = pr.presencas_por_mes(restaurant, INICIO, FIM)
        assert len(meses) == 1
        m = meses[0]
        assert (m["pessoas"], m["presencas"], m["ausencias"]) == (3, 6, Decimal(3))


class TestEcra:
    def test_ecra_e_exportacao(self, client, restaurant, manager, cenario):
        client.force_login(manager)
        r = client.get(reverse("analytics:pessoas"), {"mes": f"{MES:%Y-%m}"})
        corpo = r.content.decode()
        assert "Presenças e desvios" in corpo and "Falta de marcação" in corpo and "de Tasca da Baixa" in corpo
        assert "+" in corpo or "−" in corpo
        livro = load_workbook(io.BytesIO(client.get(reverse("analytics:pessoas"), {"mes": f"{MES:%Y-%m}", "exportar": "xlsx"}).content))
        nome = restaurant.name
        assert livro.sheetnames == [f"Pessoas {nome}", f"Áreas {nome}", f"Ausências {nome}", f"Por mês {nome}", f"Desvios {nome}"]
        linhas = list(livro[f"Desvios {nome}"].iter_rows(values_only=True))
        assert [l[2] for l in linhas[1:]] == ["Atraso", "Falta de marcação", "Saída antes da hora"]

    def test_so_admin_e_gerente(self, client, stock_user, staff_user):
        for conta in (stock_user, staff_user):
            client.force_login(conta)
            assert client.get(reverse("analytics:pessoas")).status_code == 403
        client.force_login(stock_user)
        assert reverse("analytics:pessoas") not in client.get(reverse("analytics:stock")).content.decode()
