"""Fase 15: os relatórios de stock e de compras da Análise. O Resp. stock
também os vê; o resumo, que cruza vendas e pessoas, não."""
import datetime
import io
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from analytics import stock_reports as sr
from core.models import Restaurant
from inventory.consumption import register_consumption
from inventory.models import Category, InternalConsumption, Item, StockMovement
from inventory.services import register_count, register_movement, transfer_stock
from purchasing import services as compras
from purchasing.models import Supplier

pytestmark = pytest.mark.django_db
User = get_user_model()
HOJE = timezone.localdate


def _artigo(casa, nome, categoria="Mercearia", stock="0", custo=None, user=None, unidade=Item.Unit.KG):
    cat = Category.objects.get_or_create(restaurant=casa, name=categoria)[0]
    item = Item.objects.create(restaurant=casa, name=nome, category=cat, unit=unidade,
                               reference_cost=None if custo is None else Decimal(custo))
    if Decimal(stock):
        register_movement(item=item, movement_type=StockMovement.Type.ENTRADA, quantity=stock, user=user)
        item.refresh_from_db()
    return item


def _receber(casa, fornecedor, user, dia, linhas, numero="FT 1"):
    return compras.receive(restaurant=casa, supplier=fornecedor, user=user, document_number=numero, document_date=dia,
                           lines=[(i, Decimal(q), Decimal(p), None) for i, q, p in linhas])


@pytest.fixture
def casa_b(restaurant):
    return Restaurant.objects.create(organization=restaurant.organization, name="Tasca da Baixa", slug="tasca-da-baixa")


@pytest.fixture
def dono(restaurant):
    return User.objects.create_user(email="dono@tasca.pt", password="segredo-de-teste-123",
                                    organization=restaurant.organization, is_group_admin=True)


class TestStock:
    def test_valor_por_categoria_e_por_casa(self, restaurant, casa_b, stock_user):
        _artigo(restaurant, "Arroz", stock="10", custo="1.10", user=stock_user)
        _artigo(restaurant, "Vinho", categoria="Bebidas", stock="6", custo="3.00", user=stock_user)
        _artigo(restaurant, "Sal", stock="2", user=stock_user)  # sem custo
        _artigo(casa_b, "Arroz", stock="4", custo="1.20", user=stock_user)
        linhas, totais, sem_custo = sr.valor_por_categoria([restaurant, casa_b])
        assert linhas == [
            {"categoria": "Bebidas", "valores": [Decimal("18.00"), Decimal("0.00")]},
            {"categoria": "Mercearia", "valores": [Decimal("11.00"), Decimal("4.80")]},
        ]
        assert totais == [Decimal("29.00"), Decimal("4.80")] and sem_custo == [1, 0]

    def test_mais_movimentados_por_origem(self, restaurant, casa_b, stock_user):
        arroz = _artigo(restaurant, "Arroz", stock="20", custo="1.00", user=stock_user)
        _artigo(restaurant, "Sal", stock="1", user=stock_user)
        register_movement(item=arroz, movement_type=StockMovement.Type.SAIDA, quantity=1, user=stock_user)  # quebra
        register_consumption(restaurant=restaurant, worker=None, reason=InternalConsumption.Reason.PROVA, day=HOJE(),
                             lines=[(arroz, Decimal(2))], user=stock_user)
        transfer_stock(source_item=arroz, destination_item=_artigo(casa_b, "Arroz"), quantity=3, user=stock_user)
        register_count(item=arroz, counted=Decimal("13.5"), user=stock_user)
        lista = sr.mais_movimentados(restaurant, HOJE(), HOJE())
        assert lista[0].item == arroz
        m = lista[0]
        assert (m.movimentos, m.entradas, m.quebras, m.consumos, m.transferencias, m.ajustes) == (
            5, Decimal(20), Decimal(1), Decimal(2), Decimal(-3), Decimal("-0.5"))
        assert m.valor_saidas == Decimal("3.00")

    def test_ajustes_com_valor(self, restaurant, stock_user):
        arroz = _artigo(restaurant, "Arroz", stock="10", custo="2.00", user=stock_user)
        register_count(item=arroz, counted=Decimal("9"), user=stock_user)
        register_movement(item=arroz, movement_type=StockMovement.Type.AJUSTE_POS, quantity=Decimal("0.5"), user=stock_user)
        linhas, perdas, ganhos = sr.ajustes(restaurant, HOJE(), HOJE())
        assert len(linhas) == 2 and perdas == Decimal("-2.00") and ganhos == Decimal("1.00")

    def test_ecra_e_exportacao(self, client, restaurant, stock_user):
        _artigo(restaurant, "Arroz", stock="10", custo="1.10", user=stock_user)
        client.force_login(stock_user)
        corpo = client.get(reverse("analytics:stock")).content.decode()
        assert "Valor do stock por categoria" in corpo and "11,00 €" in corpo and "Mais movimentados" in corpo
        assert "Resumo" not in corpo  # o separador do resumo não aparece ao Resp. stock
        livro = load_workbook(io.BytesIO(client.get(reverse("analytics:stock"), {"exportar": "xlsx"}).content))
        assert livro.sheetnames[0] == "Valor por categoria"
        assert list(livro["Valor por categoria"].iter_rows(values_only=True))[1] == ("Mercearia", 11.0)


class TestCompras:
    @pytest.fixture
    def dados(self, restaurant, casa_b, stock_user):
        frescos = Supplier.objects.create(organization=restaurant.organization, name="Frescos do Mercado")
        mercearia = Supplier.objects.create(organization=restaurant.organization, name="Mercearia Central")
        tomate_a, tomate_b = _artigo(restaurant, "Tomate"), _artigo(casa_b, "Tomate")
        arroz = _artigo(restaurant, "Arroz")
        dia = HOJE() - datetime.timedelta(days=20)
        _receber(restaurant, frescos, stock_user, dia, [(tomate_a, 10, "1.00")], "FT 1")
        _receber(restaurant, frescos, stock_user, dia + datetime.timedelta(days=10), [(tomate_a, 10, "1.20")], "FT 2")
        _receber(casa_b, frescos, stock_user, dia + datetime.timedelta(days=5), [(tomate_b, 5, "1.10")], "FT 3")
        _receber(restaurant, mercearia, stock_user, dia, [(arroz, 20, "1.05")], "FT 4")
        return frescos, mercearia

    def test_por_fornecedor_e_por_casa(self, restaurant, casa_b, dados):
        linhas, totais, total = sr.por_fornecedor([restaurant, casa_b], HOJE() - datetime.timedelta(days=60), HOJE())
        assert [(l["fornecedor"], l["valores"], l["total"], l["rececoes"]) for l in linhas] == [
            ("Frescos do Mercado", [Decimal("22.00"), Decimal("5.50")], Decimal("27.50"), 3),
            ("Mercearia Central", [Decimal("21.00"), Decimal("0.00")], Decimal("21.00"), 1),
        ]
        assert total == Decimal("48.50") and linhas[0]["peso"] == Decimal("56.7")

    def test_por_artigo_junta_as_casas_pelo_nome(self, restaurant, casa_b, dados):
        artigos = sr.por_artigo([restaurant, casa_b], HOJE() - datetime.timedelta(days=60), HOJE())
        tomate = next(a for a in artigos if a.nome == "Tomate")
        assert tomate.quantidade == 25 and tomate.total == Decimal("27.50")
        assert [p[1] for p in tomate.precos] == [Decimal("1.00"), Decimal("1.10"), Decimal("1.20")]
        assert (tomate.minimo, tomate.maximo, tomate.ultimo, tomate.variacao) == (
            Decimal("1.00"), Decimal("1.20"), Decimal("1.20"), Decimal("20.0"))
        assert tomate.preco_medio == Decimal("1.1000")

    def test_ecra_com_grafico_e_artigo_escolhido(self, client, restaurant, dados, stock_user):
        client.force_login(stock_user)
        corpo = client.get(reverse("analytics:compras")).content.decode()
        assert "Evolução do preço · Tomate" in corpo  # o que mais se gastou
        assert 'class="stroke-floresta"' in corpo and corpo.count("<circle") >= 4
        corpo = client.get(reverse("analytics:compras"), {"artigo": "Arroz|KG"}).content.decode()
        assert "Evolução do preço · Arroz" in corpo

    def test_por_omissao_seis_meses(self, client, restaurant, dados, stock_user):
        client.force_login(stock_user)
        r = client.get(reverse("analytics:compras"))
        assert r.context["fim"] == HOJE() and (HOJE() - r.context["inicio"]).days > 150

    def test_dono_compara_casas_e_exporta(self, client, restaurant, casa_b, dados, dono):
        client.force_login(dono)
        r = client.get(reverse("analytics:compras"), {"casa": [restaurant.pk, casa_b.pk], "exportar": "xlsx"})
        livro = load_workbook(io.BytesIO(r.content))
        assert livro.sheetnames == ["Por fornecedor", "Por artigo", "Preços Tomate"]
        cabecalho = next(livro["Por fornecedor"].iter_rows(values_only=True))
        assert cabecalho[1:3] == ("Tasca da Baixa", "Tasca do Teste")


class TestAcesso:
    def test_resp_stock_vai_do_resumo_para_o_stock(self, client, stock_user):
        client.force_login(stock_user)
        r = client.get(reverse("analytics:inicio"))
        assert r.status_code == 302 and r.url == reverse("analytics:stock")

    def test_colaborador_nao_entra(self, client, staff_user):
        client.force_login(staff_user)
        for nome in ("analytics:stock", "analytics:compras"):
            assert client.get(reverse(nome)).status_code == 403

    def test_resp_stock_so_ve_a_casa_ativa(self, client, restaurant, casa_b, stock_user):
        _artigo(casa_b, "Arroz", stock="4", custo="1.20", user=stock_user)
        client.force_login(stock_user)
        r = client.get(reverse("analytics:stock"), {"casa": [casa_b.pk]})
        assert r.context["escolhidas"] == [restaurant]
