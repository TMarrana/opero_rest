"""Fase 15: a Análise. O período e as casas lado a lado, com os mesmos
totais dos ecrãs de cada módulo, e a exportação para Excel e CSV."""
import datetime
import io
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone

from analytics import reports
from attendance.month import month_range, previous_month
from attendance.services import close_shift, open_shift
from core.middleware import SESSION_KEY
from core.models import Restaurant
from inventory.consumption import month_map, register_consumption
from inventory.models import Category, InternalConsumption, Item, StockMovement
from inventory.services import register_movement
from recipes import services as fichas
from recipes.models import Recipe
from sales import reports as vendas
from sales import sintetico
from sales.models import PosProduct
from sales.services import import_saft
from sales.sintetico import Doc
from workers.models import Worker

pytestmark = pytest.mark.django_db
User = get_user_model()
MES = previous_month(timezone.localdate())
INICIO, FIM = month_range(MES)


def _quando(dia, hora):
    return timezone.make_aware(datetime.datetime.combine(dia, datetime.time(hora)))


def _casa(restaurant, nome, slug, nif):
    casa = restaurant if nome is None else Restaurant.objects.create(
        organization=restaurant.organization, name=nome, slug=slug
    )
    casa.nif = nif
    casa.save()
    return casa


def _vendas(casa, user, docs):
    conteudo = sintetico.gerar(nif=casa.nif, documentos=docs, inicio=INICIO, fim=FIM)
    import_saft(restaurant=casa, upload=SimpleUploadedFile("saft.xml", conteudo), user=user)


def _bitoques(numero, dia, n, preco="8.94"):
    return Doc(numero, datetime.datetime.combine(dia, datetime.time(13)), [("P001", "Bitoque", n, Decimal(preco), Decimal("13"))])


def _bife(casa, user, custo="12.00"):
    categoria = Category.objects.get_or_create(restaurant=casa, name="Talho")[0]
    item = Item.objects.create(restaurant=casa, name="Bife da vazia", category=categoria, unit=Item.Unit.KG,
                               reference_cost=Decimal(custo))
    register_movement(item=item, movement_type=StockMovement.Type.ENTRADA, quantity=10, user=user)
    return item


@pytest.fixture
def grupo(restaurant, admin_user):
    """Duas casas com vendas, ponto, a mesma ficha e um consumo interno no mês
    passado."""
    a = _casa(restaurant, None, None, "509123457")
    b = _casa(restaurant, "Tasca da Baixa", "tasca-da-baixa", "510000002")
    ficha = fichas.create_recipe(
        organization=restaurant.organization, name="Bitoque", kind=Recipe.Kind.PRATO, user=admin_user,
        valid_from=INICIO, lines=[(("Bife da vazia", "KG"), Decimal("0.200"), None)],
    )
    dia = INICIO + datetime.timedelta(days=4)
    for casa, n, custo in ((a, 10, "12.00"), (b, 5, "15.00")):
        bife = _bife(casa, admin_user, custo)
        _vendas(casa, admin_user, [_bitoques(f"FS {casa.pk}/1", dia, n), Doc(f"FS {casa.pk}/2", datetime.datetime.combine(dia, datetime.time(20)), [("C001", "Café", 4, Decimal("0.70"), Decimal("23"))])])
        fichas.link_product(recipe=ficha, product=PosProduct.objects.get(restaurant=casa, code="P001"), user=admin_user)
        w = Worker.objects.create(restaurant=casa, full_name=f"Pessoa {casa.pk}")
        open_shift(worker=w, at=_quando(dia, 10))
        close_shift(worker=w, at=_quando(dia, 18))
        register_consumption(
            restaurant=casa, worker=w, reason=InternalConsumption.Reason.REFEICAO, day=dia,
            lines=[(bife, Decimal("0.200"))], user=admin_user,
        )
    return a, b


@pytest.fixture
def dono(restaurant):
    return User.objects.create_user(
        email="dono@tasca.pt", password="segredo-de-teste-123", organization=restaurant.organization, is_group_admin=True
    )


def _url(**params):
    from urllib.parse import urlencode

    return reverse("analytics:inicio") + "?" + urlencode(params, doseq=True)


class TestAceitacao:
    def test_dono_compara_as_duas_casas_e_exporta(self, client, grupo, dono):
        """O dono abre a Análise, escolhe o mês e as duas casas: vê lado a lado
        as vendas, as horas, as vendas por hora, o food cost teórico e os
        consumos internos. Exporta para Excel, e os totais batem com os de
        cada módulo."""
        a, b = grupo
        client.force_login(dono)
        r = client.get(_url(mes=f"{MES:%Y-%m}", casa=[a.pk, b.pk]))
        assert r.status_code == 200
        ind = {i.restaurant: i for i in r.context["indicadores"]}
        assert set(ind) == {a, b}
        corpo = r.content.decode()
        assert "Casas lado a lado" in corpo and "Tasca da Baixa" in corpo

        for casa in (a, b):
            i = ind[casa]
            assert i.vendas == vendas.resumo(casa, INICIO, FIM).total
            _, totais = vendas.por_dia(casa, INICIO, FIM)
            assert i.horas == totais["horas"] == datetime.timedelta(hours=8)
            assert i.por_hora == totais["por_hora"]
            mapa = month_map(casa, MES)
            assert i.consumos == sum((l["valor"] for l in mapa if l["valor"]), Decimal(0))
            assert i.refeicoes == 1

        # A: 10 bitoques a 8,94 €, 0,2 kg de bife a 12 € = 2,40 € por dose.
        assert ind[a].vendas == Decimal("92.20")
        assert ind[a].custo_teorico == Decimal("24.00")
        assert ind[a].food_cost == Decimal("26.8")
        assert ind[a].cobertura == Decimal("97.0")
        # B: a mesma ficha, mas o bife custa 15 € lá: 3,00 € por dose.
        assert ind[b].custo_teorico == Decimal("15.00")

        r = client.get(_url(mes=f"{MES:%Y-%m}", casa=[a.pk, b.pk], exportar="xlsx"))
        assert r["Content-Type"].startswith("application/vnd.openxmlformats")
        from openpyxl import load_workbook

        livro = load_workbook(io.BytesIO(r.content))
        assert livro.sheetnames[:2] == ["Casas", "Por dia"]
        folha = livro["Casas"]
        # As casas vêm pelo nome: a Baixa antes do Teste.
        assert [c.value for c in folha[1]][1:] == [b.name, a.name]
        vendas_linha = next(row for row in folha.iter_rows(values_only=True) if row[0] == "Vendas sem IVA")
        assert vendas_linha[1:] == (pytest.approx(float(ind[b].vendas)), pytest.approx(float(ind[a].vendas)))
        horas_linha = next(row for row in folha.iter_rows(values_only=True) if row[0] == "Horas trabalhadas")
        assert horas_linha[1:] == (8.0, 8.0)


class TestAcesso:
    def test_manager_sees_only_the_active_house(self, client, grupo, manager):
        a, b = grupo
        client.force_login(manager)
        r = client.get(_url(mes=f"{MES:%Y-%m}", casa=[a.pk, b.pk]))
        assert [i.restaurant for i in r.context["indicadores"]] == [a]

    def test_other_profiles_do_not_see_the_summary(self, client, grupo, stock_user, staff_user):
        """O resumo cruza vendas e pessoas: o Resp. stock vai para o stock, e o
        colaborador não entra."""
        client.force_login(stock_user)
        assert client.get(reverse("analytics:inicio")).url == reverse("analytics:stock")
        client.force_login(staff_user)
        assert client.get(reverse("analytics:inicio")).status_code == 403

    def test_group_admin_defaults_to_active_house(self, client, grupo, dono):
        a, b = grupo
        client.force_login(dono)
        sessao = client.session
        sessao[SESSION_KEY] = b.pk
        sessao.save()
        r = client.get(_url(mes=f"{MES:%Y-%m}"))
        assert [i.restaurant for i in r.context["indicadores"]] == [b]


class TestPeriodoEExportacao:
    def test_date_range(self, client, grupo, admin_user):
        a, _ = grupo
        client.force_login(admin_user)
        dia = INICIO + datetime.timedelta(days=4)
        r = client.get(_url(inicio=dia.isoformat(), fim=dia.isoformat()))
        assert r.context["inicio"] == r.context["fim"] == dia
        assert r.context["indicadores"][0].vendas == Decimal("92.20")
        r = client.get(_url(inicio=(dia + datetime.timedelta(days=1)).isoformat(), fim=FIM.isoformat()))
        assert r.context["indicadores"][0].vendas == 0

    def test_csv_export(self, client, grupo, admin_user):
        client.force_login(admin_user)
        r = client.get(_url(mes=f"{MES:%Y-%m}", exportar="csv"))
        texto = r.content.decode("utf-8")
        assert texto.startswith("﻿")
        linhas = texto.lstrip("﻿").splitlines()
        assert "Vendas sem IVA;92,20" in linhas
        assert "Horas trabalhadas;8,00" in linhas
        assert "Food cost teórico (%);26,8" in linhas

    def test_products_margin(self, grupo):
        a, _ = grupo
        produtos = {p.code: p for p in reports.produtos(a, INICIO, FIM)}
        bitoque = produtos["P001"]
        assert bitoque.cost_per_dose == Decimal("2.4") and bitoque.margin == Decimal("65.40")
        assert produtos["C001"].recipe is None and produtos["C001"].margin is None

    def test_chart_renders_one_bar_per_day(self, client, grupo, admin_user):
        client.force_login(admin_user)
        corpo = client.get(_url(mes=f"{MES:%Y-%m}")).content.decode()
        assert corpo.count("<title>") >= (FIM - INICIO).days + 1
        assert 'class="fill-floresta"' in corpo
