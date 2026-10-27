"""Fase 15: a exportação das listas em CSV e em Excel, com os mesmos filtros
do ecrã e só com os dados da casa ativa."""
import datetime
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from attendance.services import close_shift, open_shift
from core import exports
from core.models import Restaurant
from inventory.consumption import register_consumption
from inventory.models import Category, InternalConsumption, Item, StockMovement
from inventory.services import register_movement
from purchasing import services as compras
from purchasing.models import Supplier
from shifts import services as turnos
from vacations.models import AbsenceType
from vacations.services import create_request
from workers.models import Worker

pytestmark = pytest.mark.django_db
HOJE = timezone.localdate


def _csv(resp):
    assert resp["Content-Type"].startswith("text/csv")
    texto = resp.content.decode("utf-8")
    assert texto.startswith("﻿")
    return texto.lstrip("﻿").splitlines()


def _xlsx(resp):
    assert resp["Content-Type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(resp.content))


def _linhas(folha):
    return list(folha.iter_rows(values_only=True))


def _artigo(casa, nome, stock="0", custo=None, user=None, minimo="0"):
    categoria = Category.objects.get_or_create(restaurant=casa, name="Mercearia")[0]
    item = Item.objects.create(
        restaurant=casa, name=nome, category=categoria, unit=Item.Unit.KG, minimum_stock=Decimal(minimo),
        reference_cost=None if custo is None else Decimal(custo),
    )
    if Decimal(stock):
        register_movement(item=item, movement_type=StockMovement.Type.ENTRADA, quantity=stock, user=user)
        item.refresh_from_db()
    return item


@pytest.fixture
def outra_casa(restaurant):
    return Restaurant.objects.create(organization=restaurant.organization, name="Tasca da Baixa", slug="tasca-da-baixa")


class TestFormato:
    def test_csv_portugues(self):
        t = exports.Tabela(
            "T", ["Dia", "Valor", "Horas", "Quando"],
            [[datetime.date(2026, 10, 5), Decimal("1234.5"), datetime.timedelta(hours=7, minutes=30),
              timezone.make_aware(datetime.datetime(2026, 10, 5, 13, 5))]],
            formatos=[exports.DATA, exports.EUROS, exports.HORAS, exports.DATAHORA],
        )
        linhas = _csv(exports.csv_response(t, "t"))
        assert linhas == ["Dia;Valor;Horas;Quando", "05/10/2026;1234,50;7,50;05/10/2026 13:05"]

    def test_excel_numeros_e_datas(self):
        t = exports.Tabela(
            "Folha/1", ["Dia", "Valor", "Quando"],
            [[datetime.date(2026, 10, 5), Decimal("1234.5"), timezone.make_aware(datetime.datetime(2026, 10, 5, 13, 5))]],
            formatos=[exports.DATA, exports.EUROS, exports.DATAHORA],
        )
        folha = _xlsx(exports.xlsx_response([t], "t"))["Folha-1"]
        assert folha["B2"].value == 1234.5 and folha["B2"].number_format == '#,##0.00 "€"'
        assert folha["C2"].value == datetime.datetime(2026, 10, 5, 13, 5)


class TestStock:
    def test_artigos_com_filtro(self, client, restaurant, outra_casa, stock_user):
        _artigo(restaurant, "Arroz", "25", "1.10", stock_user, minimo="10")
        _artigo(restaurant, "Feijão", "2", None, stock_user, minimo="5")
        _artigo(outra_casa, "Arroz da Baixa", "1", None, stock_user)
        client.force_login(stock_user)
        linhas = _csv(client.get(reverse("inventory:lista"), {"exportar": "csv"}))
        assert linhas[0].startswith("Artigo;Categoria;Unidade;Stock")
        assert any(l.startswith("Arroz;Mercearia;quilograma;25;10;1,10;27,50;OK") for l in linhas)
        assert not any("Baixa" in l for l in linhas)
        linhas = _csv(client.get(reverse("inventory:lista"), {"exportar": "csv", "baixo": "1"}))
        assert len(linhas) == 2 and linhas[1].startswith("Feijão;")

    def test_links_no_fragmento_htmx_levam_os_filtros(self, client, restaurant, stock_user):
        _artigo(restaurant, "Arroz", "1", None, stock_user)
        client.force_login(stock_user)
        corpo = client.get(reverse("inventory:lista"), {"q": "arr"}, HTTP_HX_REQUEST="true").content.decode()
        assert "?q=arr&amp;exportar=xlsx" in corpo

    def test_movimentos_todos_sem_limite(self, client, restaurant, stock_user):
        arroz = _artigo(restaurant, "Arroz", "1000", None, stock_user)
        for _ in range(205):
            register_movement(item=arroz, movement_type=StockMovement.Type.SAIDA, quantity=1, user=stock_user)
        client.force_login(stock_user)
        folha = _xlsx(client.get(reverse("inventory:historico"), {"artigo": arroz.pk, "exportar": "xlsx"}))["Movimentos"]
        linhas = _linhas(folha)
        assert len(linhas) == 1 + 206
        assert linhas[1][4] == -1  # a saída com sinal


class TestPessoas:
    def test_ponto_do_mes_so_desta_casa(self, client, restaurant, outra_casa, manager):
        marta = Worker.objects.create(restaurant=restaurant, full_name="Marta Silva")
        ana = Worker.objects.create(restaurant=outra_casa, full_name="Ana da Baixa")
        dia = HOJE() - datetime.timedelta(days=1)
        for w in (marta, ana):
            open_shift(worker=w, at=timezone.make_aware(datetime.datetime.combine(dia, datetime.time(10))))
            close_shift(worker=w, at=timezone.make_aware(datetime.datetime.combine(dia, datetime.time(16))))
        client.force_login(manager)
        folha = _xlsx(client.get(reverse("attendance:dia"), {"dia": dia.isoformat(), "exportar": "xlsx"}))["Ponto"]
        linhas = _linhas(folha)
        assert [l[0] for l in linhas[1:]] == ["Marta Silva"]
        assert linhas[1][6] == 6.0

    def test_ausencias_com_filtro(self, client, restaurant, manager):
        marta = Worker.objects.create(restaurant=restaurant, full_name="Marta Silva")
        rui = Worker.objects.create(restaurant=restaurant, full_name="Rui Costa")
        ferias = AbsenceType.objects.get(restaurant=restaurant, name="Férias")
        for w in (marta, rui):
            create_request(worker=w, absence_type=ferias, start_date=HOJE() + datetime.timedelta(days=30),
                           end_date=HOJE() + datetime.timedelta(days=31))
        client.force_login(manager)
        linhas = _csv(client.get(reverse("vacations:lista"), {"trabalhador": rui.pk, "exportar": "csv"}))
        assert len(linhas) == 2 and linhas[1].startswith("Rui Costa;Férias;")

    def test_turnos_da_semana(self, client, restaurant, manager):
        marta = Worker.objects.create(restaurant=restaurant, full_name="Marta Silva")
        semana = turnos.week_start(HOJE()) + datetime.timedelta(days=7)
        turnos.create_shift(restaurant=restaurant, worker=marta, day=semana, start=datetime.time(12), end=datetime.time(16))
        turnos.set_rest_day(restaurant=restaurant, worker=marta, day=semana + datetime.timedelta(days=1))
        client.force_login(manager)
        linhas = _csv(client.get(reverse("shifts:planeador"), {"semana": semana.isoformat(), "exportar": "csv"}))
        assert linhas[1].startswith(f"Marta Silva;{semana:%d/%m/%Y};{semana:%d/%m/%Y} 12:00;{semana:%d/%m/%Y} 16:00;4,00;Turno")
        assert linhas[2].startswith("Marta Silva;") and ";Folga;" in linhas[2]

    def test_colaborador_nao_exporta(self, client, staff_user):
        client.force_login(staff_user)
        assert client.get(reverse("attendance:dia"), {"exportar": "csv"}).status_code == 403


class TestComprasEConsumos:
    def test_encomendas_com_linhas(self, client, restaurant, stock_user):
        fornecedor = Supplier.objects.create(organization=restaurant.organization, name="Mercearia Central")
        arroz = _artigo(restaurant, "Arroz", "0", None, stock_user)
        compras.create_order(restaurant=restaurant, supplier=fornecedor, lines=[(arroz, Decimal(10), Decimal("1.05"))], user=stock_user)
        client.force_login(stock_user)
        livro = _xlsx(client.get(reverse("purchasing:encomendas"), {"estado": "todas", "exportar": "xlsx"}))
        assert livro.sheetnames == ["Encomendas", "Linhas"]
        assert _linhas(livro["Encomendas"])[1][1] == "Mercearia Central"
        assert _linhas(livro["Linhas"])[1][2:7] == ("Arroz", "quilograma", 10, 1.05, 10.5)

    def test_rececoes_em_excel_e_csv_de_sempre(self, client, restaurant, stock_user):
        fornecedor = Supplier.objects.create(organization=restaurant.organization, name="Mercearia Central")
        arroz = _artigo(restaurant, "Arroz", "0", None, stock_user)
        compras.receive(restaurant=restaurant, supplier=fornecedor, user=stock_user, document_number="FT 9",
                        document_date=HOJE(), lines=[(arroz, Decimal(10), Decimal("1.05"), None)])
        client.force_login(stock_user)
        mes = HOJE().strftime("%Y-%m")
        folha = _xlsx(client.get(reverse("purchasing:rececoes"), {"mes": mes, "exportar": "xlsx"}))["Receções"]
        assert _linhas(folha)[1][6:] == ("Arroz", 10, "KG", 1.05, 10.5)
        texto = client.get(reverse("purchasing:rececoes"), {"mes": mes, "exportar": "1"}).content.decode("utf-8-sig")
        assert ";Arroz;10;KG;1,05;10,5" in texto

    def test_consumos_lista_e_mapa(self, client, restaurant, manager):
        marta = Worker.objects.create(restaurant=restaurant, full_name="Marta Silva")
        cafe = _artigo(restaurant, "Café", "10", "0.20", manager)
        register_consumption(restaurant=restaurant, worker=marta, reason=InternalConsumption.Reason.OFERTA, day=HOJE(),
                             lines=[(cafe, Decimal(2))], user=manager)
        client.force_login(manager)
        mes = HOJE().strftime("%Y-%m")
        livro = _xlsx(client.get(reverse("inventory:consumos"), {"mes": mes, "exportar": "xlsx"}))
        assert livro.sheetnames == ["Consumos", "Mapa do mês"]
        assert _linhas(livro["Consumos"])[1][2:9] == ("Marta Silva", "Oferta", "Café", 2, "KG", 0.2, 0.4)
        # O mapa para o contabilista continua igual.
        assert client.get(reverse("inventory:consumos"), {"mes": mes, "exportar": "1"})["Content-Type"].startswith("text/csv")


class TestVendas:
    def test_excel_por_dia_e_por_produto(self, client, restaurant, manager):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from sales import sintetico
        from sales.services import import_saft
        from sales.sintetico import Doc

        restaurant.nif = "509123457"
        restaurant.save()
        ontem = HOJE() - datetime.timedelta(days=1)
        doc = Doc("FS E/1", datetime.datetime.combine(ontem, datetime.time(13)), [("P001", "Bitoque", 2, Decimal("8.94"), Decimal("13"))])
        conteudo = sintetico.gerar(nif=restaurant.nif, documentos=[doc], inicio=ontem, fim=ontem)
        import_saft(restaurant=restaurant, upload=SimpleUploadedFile("s.xml", conteudo), user=manager)
        client.force_login(manager)
        params = {"de": ontem.isoformat(), "ate": ontem.isoformat()}
        folha = _xlsx(client.get(reverse("sales:vendas"), {**params, "exportar": "dias", "formato": "xlsx"}))["Vendas por dia"]
        assert _linhas(folha)[1][1] == pytest.approx(17.88)
        folha = _xlsx(client.get(reverse("sales:vendas"), {**params, "exportar": "produtos", "formato": "xlsx"}))["Vendas por produto"]
        assert _linhas(folha)[1][:5] == ("P001", "Bitoque", "Pratos", 2, 17.88)
