"""
Exportação de tabelas para CSV e Excel (Fase 15).

O CSV segue o das outras exportações: `;`, BOM e vírgula decimal, como o
Excel em português espera. No Excel (`openpyxl`), os números ficam números,
com o formato da coluna, para se poder somar.
"""
import csv
import datetime
import io
from dataclasses import dataclass, field
from decimal import Decimal

from django.http import HttpResponse

TEXTO, EUROS, NUMERO, PERCENTAGEM, HORAS, DATA, DATAHORA = "texto", "euros", "numero", "pct", "horas", "data", "datahora"
FORMATOS = ("csv", "xlsx")

_FORMATO_XLSX = {
    EUROS: '#,##0.00 "€"',
    NUMERO: "#,##0.###",
    PERCENTAGEM: '0.0 "%"',
    HORAS: "0.00",
    DATA: "dd/mm/yyyy",
    DATAHORA: "dd/mm/yyyy hh:mm",
}


@dataclass
class Tabela:
    """Uma tabela para exportar. `formatos` diz, coluna a coluna, como se
    escreve cada valor; sem ele, tudo é texto. Numa tabela em que cada linha
    é uma medida (as casas lado a lado), `formatos_linha` diz o formato de
    cada linha, a partir da segunda coluna."""

    titulo: str
    colunas: list
    linhas: list = field(default_factory=list)
    formatos: list = field(default_factory=list)
    formatos_linha: list = field(default_factory=list)

    def formato(self, i, linha=None):
        if self.formatos_linha and linha is not None and i > 0:
            return self.formatos_linha[linha]
        return self.formatos[i] if i < len(self.formatos) else TEXTO


def em_horas(delta):
    """Um timedelta em horas decimais, para exportar."""
    if delta is None:
        return None
    return (Decimal(delta.total_seconds()) / 3600).quantize(Decimal("0.01"))


def _texto_csv(valor, formato):
    if valor is None:
        return ""
    if isinstance(valor, datetime.timedelta):
        valor = em_horas(valor)
    if isinstance(valor, datetime.datetime):
        return _local(valor).strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, datetime.date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "sim" if valor else ""
    if isinstance(valor, Decimal):
        casas = 2 if formato in (EUROS, HORAS) else (1 if formato == PERCENTAGEM else None)
        valor = valor.quantize(Decimal(10) ** -casas) if casas is not None else valor.normalize()
        return format(valor, "f").replace(".", ",")
    return str(valor)


def csv_response(tabela, nome):
    saida = io.StringIO()
    escrita = csv.writer(saida, delimiter=";")
    escrita.writerow(tabela.colunas)
    for n, linha in enumerate(tabela.linhas):
        escrita.writerow([_texto_csv(v, tabela.formato(i, n)) for i, v in enumerate(linha)])
    resp = HttpResponse("﻿" + saida.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{nome}.csv"'
    return resp


def _local(valor):
    """Hora local e sem fuso: o Excel não guarda fusos horários."""
    from django.utils import timezone

    if timezone.is_aware(valor):
        valor = timezone.localtime(valor)
    return valor.replace(tzinfo=None)


def _valor_xlsx(valor):
    if isinstance(valor, datetime.datetime):
        return _local(valor)
    if isinstance(valor, bool):
        return "sim" if valor else ""
    if isinstance(valor, datetime.timedelta):
        return float(em_horas(valor))
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def xlsx_response(tabelas, nome):
    """Um livro com uma folha por tabela."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    livro.remove(livro.active)
    for tabela in tabelas:
        # O Excel não aceita mais de 31 caracteres, nem / no nome da folha.
        folha = livro.create_sheet(tabela.titulo.replace("/", "-")[:31])
        folha.append(tabela.colunas)
        for celula in folha[1]:
            celula.font = Font(bold=True)
        for linha in tabela.linhas:
            folha.append([_valor_xlsx(v) for v in linha])
        for i, coluna in enumerate(tabela.colunas):
            letra = get_column_letter(i + 1)
            for n, (celula,) in enumerate(folha.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1)):
                formato = _FORMATO_XLSX.get(tabela.formato(i, n))
                if formato:
                    celula.number_format = formato
            largura = max([len(str(coluna))] + [len(str(l[i])) for l in tabela.linhas if i < len(l) and l[i] is not None])
            folha.column_dimensions[letra].width = min(max(largura + 2, 10), 50)
        folha.freeze_panes = "A2"
    saida = io.BytesIO()
    livro.save(saida)
    resp = HttpResponse(
        saida.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="{nome}.xlsx"'
    return resp


def formato_pedido(request, chave="exportar"):
    """O formato pedido no endereço (`?exportar=csv` ou `xlsx`), ou None."""
    valor = request.GET.get(chave)
    return valor if valor in FORMATOS else None


def responder(formato, tabelas, nome):
    """O CSV leva a primeira tabela; o Excel, uma folha por tabela."""
    if formato == "csv":
        return csv_response(tabelas[0], nome)
    return xlsx_response(tabelas, nome)
